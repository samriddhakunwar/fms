# pyright: reportAttributeAccessIssue=false

from datetime import timedelta
from decimal import Decimal

from django.urls import reverse
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import User
from inventory.models import Product
from sales.models import Sale, SaleItem


class SalesReportTests(APITestCase):
    def setUp(self):
        self.password = "TestPass123!"
        self.admin = User.objects.create_user(
            username="admin_user", password=self.password, role=User.Role.ADMIN
        )
        self.manager = User.objects.create_user(
            username="manager_user",
            password=self.password,
            role=User.Role.MANAGER,
        )
        self.employee = User.objects.create_user(
            username="employee_user",
            password=self.password,
            role=User.Role.STAFF,
        )

        self.chair = Product.objects.create(
            product_name="Chair",
            sku="SKU-CHAIR",
            selling_price=Decimal("500.00"),
            quantity_in_stock=500,
            minimum_stock_level=10,
        )
        self.table = Product.objects.create(
            product_name="Table",
            sku="SKU-TABLE",
            selling_price=Decimal("2000.00"),
            quantity_in_stock=500,
            minimum_stock_level=2,
        )

        self.today = timezone.localdate()
        self._sale(self.today, [(self.chair, 2)])
        self._sale(self.today, [(self.table, 1)])
        self._sale(self.today - timedelta(days=2), [(self.chair, 1)])

    def _sale(self, day, lines):
        when = timezone.make_aware(
            timezone.datetime.combine(day, timezone.datetime.min.time().replace(hour=12))
        )
        sale = Sale.objects.create(
            invoice_number=f"INV-T{Sale.objects.count():04d}",
            sold_to="Shyam Pvt. Ltd.",
            sale_date=when,
        )
        total = Decimal("0")
        for product, quantity in lines:
            item = SaleItem(
                sale=sale,
                product=product,
                quantity=quantity,
                unit_price=product.selling_price,
            )
            item.save()
            total += item.subtotal
        sale.total_amount = total
        sale.save(update_fields=["total_amount"])
        return sale

    # Access
    def test_admin_can_generate_the_report(self):
        self.client.login(username="admin_user", password=self.password)
        response = self.client.get(reverse("report_sales"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_manager_can_generate_the_report(self):
        self.client.login(username="manager_user", password=self.password)
        response = self.client.get(reverse("report_sales"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_employee_cannot_generate_the_report(self):
        self.client.login(username="employee_user", password=self.password)
        response = self.client.get(reverse("report_sales"))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_anonymous_cannot_generate_the_report(self):
        response = self.client.get(reverse("report_sales"))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    # Content
    def test_default_window_is_the_last_seven_days(self):
        self.client.login(username="admin_user", password=self.password)
        data = self.client.get(reverse("report_sales")).data

        self.assertEqual(len(data["days_series"]), 7)
        self.assertEqual(data["end_date"], self.today.isoformat())
        self.assertEqual(data["totals"]["sales_count"], 3)
        self.assertEqual(
            Decimal(data["totals"]["revenue"]), Decimal("1000.00") + Decimal("2000.00") + Decimal("500.00")
        )
        self.assertEqual(data["totals"]["items_sold"], 4)

    def test_report_breaks_revenue_down_by_product(self):
        self.client.login(username="admin_user", password=self.password)
        products = self.client.get(reverse("report_sales")).data["products"]

        by_sku = {entry["sku"]: entry for entry in products}
        self.assertEqual(by_sku["SKU-CHAIR"]["quantity"], 3)
        self.assertEqual(Decimal(by_sku["SKU-CHAIR"]["revenue"]), Decimal("1500.00"))
        self.assertEqual(by_sku["SKU-TABLE"]["quantity"], 1)
        # Ordered by quantity sold, highest first.
        self.assertEqual(products[0]["sku"], "SKU-CHAIR")

    def test_extremes_ignore_days_without_sales(self):
        self.client.login(username="admin_user", password=self.password)
        data = self.client.get(reverse("report_sales")).data

        self.assertEqual(data["busiest_day"]["day"], self.today.isoformat())
        self.assertEqual(data["busiest_day"]["count"], 2)
        self.assertEqual(
            data["quietest_day"]["day"], (self.today - timedelta(days=2)).isoformat()
        )
        self.assertEqual(data["quietest_day"]["count"], 1)
        self.assertEqual(data["best_revenue_day"]["day"], self.today.isoformat())
        self.assertEqual(data["totals"]["selling_days"], 2)

    def test_series_is_dense_so_quiet_days_report_zero(self):
        self.client.login(username="admin_user", password=self.password)
        series = self.client.get(reverse("report_sales")).data["days_series"]

        yesterday = (self.today - timedelta(days=1)).isoformat()
        quiet = next(day for day in series if day["day"] == yesterday)
        self.assertEqual(quiet["count"], 0)
        self.assertEqual(Decimal(quiet["revenue"]), Decimal("0"))

    def test_date_range_filters_the_report(self):
        self.client.login(username="admin_user", password=self.password)
        response = self.client.get(
            reverse("report_sales"),
            {"start_date": self.today.isoformat(), "end_date": self.today.isoformat()},
        )
        data = response.data

        self.assertEqual(len(data["days_series"]), 1)
        self.assertEqual(data["totals"]["sales_count"], 2)
        self.assertEqual(Decimal(data["totals"]["revenue"]), Decimal("3000.00"))

    def test_empty_range_reports_zeros_rather_than_failing(self):
        self.client.login(username="admin_user", password=self.password)
        far_off = (self.today - timedelta(days=90)).isoformat()
        data = self.client.get(
            reverse("report_sales"), {"start_date": far_off, "end_date": far_off}
        ).data

        self.assertEqual(data["totals"]["sales_count"], 0)
        self.assertEqual(Decimal(data["totals"]["revenue"]), Decimal("0"))
        self.assertIsNone(data["busiest_day"])
        self.assertEqual(data["products"], [])

    def test_backwards_range_is_rejected(self):
        self.client.login(username="admin_user", password=self.password)
        response = self.client.get(
            reverse("report_sales"),
            {
                "start_date": self.today.isoformat(),
                "end_date": (self.today - timedelta(days=5)).isoformat(),
            },
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_malformed_date_is_rejected(self):
        self.client.login(username="admin_user", password=self.password)
        response = self.client.get(reverse("report_sales"), {"start_date": "last-tuesday"})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_oversized_range_is_clamped_and_says_so(self):
        self.client.login(username="admin_user", password=self.password)
        data = self.client.get(
            reverse("report_sales"),
            {
                "start_date": (self.today - timedelta(days=400)).isoformat(),
                "end_date": self.today.isoformat(),
            },
        ).data

        self.assertTrue(data["truncated"])
        self.assertEqual(len(data["days_series"]), 180)


class ReportRecordsTests(SalesReportTests):
    """The tabular preview and its PDF / Excel downloads."""

    def _get(self, start, end, user="admin_user", **extra):
        self.client.login(username=user, password=self.password)
        return self.client.get(
            reverse("report_sales_records"),
            {"start_date": start.isoformat(), "end_date": end.isoformat(), **extra},
        )

    def test_sales_rows_cover_the_range_including_the_end_date(self):
        data = self._get(self.today - timedelta(days=1), self.today).data
        self.assertEqual(len(data["rows"]), 2)  # the sale two days ago is excluded
        summary = {entry["label"]: entry["value"] for entry in data["summary"]}
        self.assertEqual(summary["Invoices"], 2)
        self.assertEqual(Decimal(summary["Total revenue"]), Decimal("3000.00"))

    def test_sale_late_on_the_end_date_is_included(self):
        sale = self._sale(self.today - timedelta(days=5), [(self.chair, 1)])
        sale.sale_date = sale.sale_date.replace(hour=23, minute=59, second=59)
        sale.save(update_fields=["sale_date"])
        day = self.today - timedelta(days=5)
        self.assertEqual(len(self._get(day, day).data["rows"]), 1)

    def test_empty_range_returns_no_rows(self):
        far_off = self.today - timedelta(days=90)
        self.assertEqual(self._get(far_off, far_off).data["rows"], [])

    def test_bad_ranges_are_rejected(self):
        self.client.login(username="admin_user", password=self.password)
        url = reverse("report_sales_records")
        for params in (
            {},
            {"start_date": "2026-13-01", "end_date": "2026-10-12"},
            {"start_date": "2026-10-12", "end_date": "2026-10-01"},
            {"start_date": "2025-01-01", "end_date": "2026-10-12"},
            {"start_date": "2026-10-01", "end_date": "2026-10-12", "export": "csv"},
        ):
            self.assertEqual(self.client.get(url, params).status_code, 400, params)

    def test_manager_can_read_the_rows(self):
        response = self._get(self.today, self.today, user="manager_user")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_staff_and_anonymous_cannot_read_or_download(self):
        response = self._get(self.today, self.today, user="employee_user", export="pdf")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.client.logout()
        response = self.client.get(reverse("report_sales_records"))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_downloads_match_the_preview(self):
        from io import BytesIO

        from openpyxl import load_workbook

        start = self.today - timedelta(days=1)
        preview = self._get(start, self.today).data
        stock_before = list(Product.objects.values_list("quantity_in_stock", flat=True))

        response = self._get(start, self.today, export="xlsx")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(
            f"Sales_Report_{start}_to_{self.today}.xlsx", response["Content-Disposition"]
        )
        sheet = load_workbook(BytesIO(response.content)).active
        invoices = [row[1] for row in sheet.iter_rows(values_only=True) if str(row[1]).startswith("INV-")]
        self.assertEqual(invoices, [row[1] for row in preview["rows"]])

        response = self._get(start, self.today, export="pdf")
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

        # Generating reports never touches stock.
        self.assertEqual(
            list(Product.objects.values_list("quantity_in_stock", flat=True)), stock_before
        )
